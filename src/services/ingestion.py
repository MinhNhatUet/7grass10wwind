"""Trích xuất PDF/PPTX và chia chunk (BL-10; FR-06.2, FR-07.2–07.4; PRD 9.3, 10.2).

- Loại file xác định bằng nội dung (magic bytes), không dựa đuôi file (FR-06.2).
- Mỗi chunk nằm trọn trong một trang/slide để locator luôn đúng; trang dài được chia nhỏ có chồng lấn.
- Trang/slide rỗng hoặc gần rỗng được đánh dấu theo vị trí, không điền nội dung suy đoán (FR-07.4).
- Không OCR (Should BL-33): PDF scan sẽ vào NeedsReview.
- Text được làm sạch theo từng trang trước khi chia chunk (src/services/text_cleaning.py, D-29).
- Phân loại từng trang PDF: trang scan (không có text dùng được nhưng có ảnh) và trang nhiều công thức được liệt kê
  trong Extraction.enrich để job nền page_enrichment OCR/kiểm tra bằng vision LLM. Trích xuất nhanh không bị chậm lại.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

from src.config import get_settings
from src.ingestion.quality.text_quality import evaluate_native_text
from src.services.text_cleaning import SUB, SUP, clean_pages

MIN_TEXT_CHARS = 40  # dưới ngưỡng này coi là trang gần rỗng (đề xuất)
SCAN_CORRUPTION_RATIO = 0.2  # lớp text của bản scan bị hỏng quá ngưỡng này → coi như không có text (đề xuất)

BIG_OPS = set("∑∏∫∮⋃⋂⨁⨂")  # toán tử lớn có baseline riêng — không dùng làm baseline dòng
# Ký hiệu đánh dấu trang có công thức: toán tử, quan hệ, logic, chữ Hy Lạp và chỉ số dưới/trên dựng lại từ pdfium
SCRIPT_CHARS = set(SUB.values()) | set(SUP.values())
SCRIPT_RUN = re.compile(f"[{re.escape(''.join(sorted(SCRIPT_CHARS)))}]+")
MATH_SYMBOLS = (
    set("∑∏∫∮√∞≤≥≠≈≡∼≅∝∈∉∋⊂⊆⊃⊇∪∩∅∀∃¬∧∨⊕⊗→←↔⇒⇐⇔↦∂∇±∓×÷∘⌊⌋⌈⌉′‖⊥")
    | set("αβγδεζηθικλμνξπρστυφχψωΓΔΘΛΞΠΣΦΨΩ")
    | SCRIPT_CHARS
)


def _font_name(pdfium_c, raw, index: int) -> str:  # noqa: ANN001
    import ctypes

    buf = ctypes.create_string_buffer(128)
    flags = ctypes.c_int()
    size = pdfium_c.FPDFText_GetFontInfo(raw, index, buf, 128, flags)
    return buf.raw[: max(size - 1, 0)].decode("latin-1", "ignore")


def _pdfium_page_text(textpage) -> str:  # noqa: ANN001 — pypdfium2.PdfTextPage
    """Text trang từ pdfium, chữ cỡ nhỏ lệch baseline → chỉ số dưới/trên; bỏ xuống dòng sinh ra giữa công thức."""
    import ctypes
    import statistics

    import pypdfium2.raw as pdfium_c

    n = textpage.count_chars()
    raw = textpage.raw
    sizes = [pdfium_c.FPDFText_GetFontSize(raw, i) for i in range(n)]
    real = [s for s in sizes if s > 1.5]  # ký tự khoảng trắng pdfium tự sinh có cỡ 1
    if not real:
        return textpage.get_text_range()
    body = statistics.median(real)
    x, y = ctypes.c_double(), ctypes.c_double()
    out: list[str] = []
    kinds: list[str] = []  # "n" thường, "s" dưới, "p" trên, "b" xuống dòng
    base: float | None = None
    for i in range(n):
        ch = chr(pdfium_c.FPDFText_GetUnicode(raw, i))
        if ch in "\r\n":
            out.append(ch)
            kinds.append("b")
            continue
        pdfium_c.FPDFText_GetCharOrigin(raw, i, x, y)
        kind = "n"
        if 1.5 < sizes[i] < 0.85 * body and base is not None and abs(y.value - base) < 0.8 * body:
            if y.value - base > 0.5:
                kind = "p"
            elif y.value - base < -0.5:
                kind = "s"
        elif sizes[i] > 1.5 and ch not in BIG_OPS and not ch.isspace():
            base = y.value
        if kind == "p" and ch == "0" and "CMSY" in _font_name(pdfium_c, raw, i):
            ch = "′"  # dấu phẩy trên (prime) của TeX nằm ở mã "0" trong font CMSY: a′, f′(x)
        out.append((SUP if kind == "p" else SUB).get(ch, ch) if kind in "ps" else ch)
        kinds.append(kind)
    # Xuống dòng nằm sát chỉ số (giới hạn của ∑ bị trích thành dòng riêng) → bỏ
    text = []
    for i, ch in enumerate(out):
        if kinds[i] == "b":
            prev = next((kinds[j] for j in range(i - 1, -1, -1) if kinds[j] != "b"), "n")
            nxt = next((kinds[j] for j in range(i + 1, len(kinds)) if kinds[j] != "b"), "n")
            if prev in "ps" or nxt in "ps":
                continue
        text.append(ch)
    return "".join(text)


class IngestionError(Exception):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass
class ExtractedChunk:
    locator_type: str
    locator_value: int
    text: str
    quality: str


@dataclass
class Extraction:
    file_type: str
    page_count: int
    chunks: list[ExtractedChunk]
    warnings: list[dict] = field(default_factory=list)
    cleaning: dict = field(default_factory=dict)  # thống kê làm sạch (ghi vào audit)
    enrich: list[dict] = field(default_factory=list)  # [{"locator_value": n, "kind": "scan"|"formula"}]


def math_symbol_count(text: str) -> int:
    """Số ký hiệu toán; mỗi cụm chỉ số dưới/trên tính là một. Bỏ qua "từ" có chỉ số kiểu rác do chữ nhỏ bị dựng nhầm
    thành chỉ số (₍McGᵣₐw₋Hᵢₗₗ, ₂₀₀₉₎): ≥3 chữ cái chỉ số hoặc ≥4 chữ số chỉ số trong một từ — chỉ số thật thì ngắn."""
    count = 0
    for token in text.split():
        count += sum(1 for ch in token if ch in MATH_SYMBOLS and ch not in SCRIPT_CHARS)
        runs = SCRIPT_RUN.findall(token)
        letters = sum(1 for ch in token if ch in SCRIPT_CHARS and ch.isalpha())
        digits = sum(1 for ch in token if ch in SCRIPT_CHARS and ch.isnumeric())
        if runs and letters < 3 and digits < 4:
            count += len(runs)
    return count


def classify_page(cleaned: str, has_image: bool, raw: str | None = None) -> str:
    """'scan' | 'formula' | 'text' — chỉ dựa trên text pdfium và việc trang có ảnh, không tốn thêm lượt đọc file.

    Đếm ký hiệu trên text ĐÃ làm sạch: header/footer lặp (vd. dòng chữ nhỏ bị dựng nhầm thành chỉ số) không làm
    trang bị coi là trang công thức. Tỷ lệ ký tự hỏng đo trên text thô vì bước làm sạch đã bỏ ký tự điều khiển.
    """
    text = _normalize(cleaned)
    corrupted = evaluate_native_text(raw).corruption_ratio > SCAN_CORRUPTION_RATIO if raw and raw.strip() else False
    unreadable = len(text) < MIN_TEXT_CHARS or corrupted
    if unreadable and has_image:
        return "scan"
    if math_symbol_count(text) >= get_settings().enrich_formula_min_symbols:
        return "formula"
    return "text"


def detect_file_type(content: bytes) -> str:
    if content.startswith(b"%PDF-"):
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile as e:
            raise IngestionError("CORRUPT_FILE") from e
        if "ppt/presentation.xml" in names:
            if any(n.endswith("vbaProject.bin") for n in names):
                raise IngestionError("MACRO_NOT_ALLOWED")  # FR-06: không thực thi nội dung nhúng/macro
            return "pptx"
    raise IngestionError("UNSUPPORTED_FILE_TYPE")


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split(text: str, max_chars: int, overlap: int) -> list[str]:
    """Chia theo đoạn/câu, không vượt max_chars; chồng lấn overlap ký tự giữa các phần."""
    if len(text) <= max_chars:
        return [text]
    parts, start = [], 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            window = text[start:end]
            cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("\n"))
            if cut > max_chars // 2:
                end = start + cut + 1
        parts.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
        # phần chồng lấn bắt đầu ở ranh giới từ, không cắt giữa từ ("Theorem" → "heorem")
        while 0 < start < end and not text[start - 1].isspace():
            start += 1
    return [p for p in parts if p]


def _page_chunks(
    locator_type: str, number: int, raw: str, quality: str = "ok"
) -> tuple[list[ExtractedChunk], dict | None]:
    s = get_settings()
    text = _normalize(raw)
    if len(text) < MIN_TEXT_CHARS:
        warning = {"locator_type": locator_type, "locator_value": number, "issue": "empty_or_no_text"}
        if not text:
            return [], warning
        return [ExtractedChunk(locator_type, number, text, "low")], warning
    return [
        ExtractedChunk(locator_type, number, p, quality) for p in _split(text, s.chunk_max_chars, s.chunk_overlap_chars)
    ], None


def page_chunks(locator_type: str, number: int, text: str, quality: str) -> list[ExtractedChunk]:
    """Chia chunk một trang đã được OCR/bổ sung (page_enrichment) theo đúng quy tắc của trích xuất nhanh."""
    return _page_chunks(locator_type, number, text, quality)[0]


def extract_pdf(content: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise IngestionError("ENCRYPTED_FILE")
        pages = reader.pages
        count = len(pages)
    except PdfReadError as e:
        raise IngestionError("CORRUPT_FILE") from e
    if count > get_settings().upload_max_pages:
        raise IngestionError("TOO_MANY_PAGES")
    # Text lấy bằng pdfium: đọc đúng font toán của TeX (≥, ∪, ·, dấu phẩy); pypdf trả ký tự rác (\x15, ";") → mất
    # ký hiệu trong trích dẫn (D-41). Trang pdfium lỗi → dùng pypdf cho trang đó.
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(content)
    except Exception:  # noqa: BLE001
        pdf = None
    raw_pages = []
    images: dict[int, bool] = {}
    for i, page in enumerate(pages, start=1):
        raw = ""
        has_image = False
        if pdf is not None:
            try:
                pdf_page = pdf[i - 1]
                textpage = pdf_page.get_textpage()
                try:
                    raw = _pdfium_page_text(textpage)
                    has_image = _has_image(pdf_page)
                finally:
                    textpage.close()
                    pdf_page.close()
            except Exception:  # noqa: BLE001
                raw = ""
        if not raw.strip():
            try:
                raw = page.extract_text() or ""
            except Exception:  # noqa: BLE001 — một trang lỗi không làm hỏng cả file; đánh dấu vị trí
                raw = ""
        raw_pages.append(("page", i, raw))
        images[i] = has_image
    if pdf is not None:
        pdf.close()
    return _build("pdf", count, raw_pages, images)


def _has_image(pdf_page) -> bool:  # noqa: ANN001 — pypdfium2.PdfPage
    import pypdfium2.raw as pdfium_c

    return next(iter(pdf_page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE], max_depth=2)), None) is not None


def extract_pptx(content: bytes) -> Extraction:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001
        raise IngestionError("CORRUPT_FILE") from e
    slides = list(prs.slides)
    if len(slides) > get_settings().upload_max_pages:
        raise IngestionError("TOO_MANY_PAGES")
    raw_pages = []
    for i, slide in enumerate(slides, start=1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text for cell in row.cells))
        raw_pages.append(("slide", i, "\n".join(texts)))
    return _build("pptx", len(slides), raw_pages)


def _build(
    file_type: str, count: int, raw_pages: list[tuple[str, int, str]], images: dict[int, bool] | None = None
) -> Extraction:
    """Làm sạch toàn tài liệu (cần thống kê header lặp giữa các trang) rồi chia chunk theo từng trang."""
    s = get_settings()
    cleaned, stats = clean_pages(raw_pages)
    raw_by_number = {number: raw for _, number, raw in raw_pages}
    toc = set(stats.toc_pages) | set(stats.index_pages)
    chunks, warnings, enrich = [], [], []
    for loc_type, number, text in cleaned:
        if number in toc:
            continue  # trang mục lục/chỉ mục: không phải nội dung học, không index
        cs, w = _page_chunks(loc_type, number, text)
        chunks += cs
        if w:
            warnings.append(w)
        if images is None or not s.enrich_enabled:  # PPTX: chưa hỗ trợ bổ sung
            continue
        kind = classify_page(text, images.get(number, False), raw_by_number.get(number))
        if kind != "text":
            if len(enrich) < s.enrich_max_pages:
                enrich.append({"locator_value": number, "kind": kind})
            else:  # vượt trần chi phí: không OCR/kiểm tra, để người duyệt xem trang này
                warnings.append({"locator_type": loc_type, "locator_value": number, "issue": f"{kind}_not_enriched"})
    return Extraction(file_type, count, chunks, warnings, stats.as_dict(), enrich)


def extract(content: bytes) -> Extraction:
    file_type = detect_file_type(content)
    return extract_pdf(content) if file_type == "pdf" else extract_pptx(content)


def needs_review(extraction: Extraction) -> bool:
    """Không có text dùng được hoặc hơn 20% trang lỗi → NeedsReview (ngưỡng đề xuất)."""
    if not extraction.chunks:
        return True
    return len(extraction.warnings) > 0.2 * max(extraction.page_count, 1)
