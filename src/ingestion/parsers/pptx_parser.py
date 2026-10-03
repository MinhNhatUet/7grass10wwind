from __future__ import annotations

from pathlib import Path

from src.ingestion.models.element import Element, ElementType
from src.ingestion.models.page import Locator, PageIR, ParsedDocument
from src.ingestion.parsers.base import BaseParser, ParserDependencyError
from src.ingestion.quality.detector import detect_page_quality


class PPTXNativeParser(BaseParser):
    def can_handle(self, mime_type: str) -> bool:
        return mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def parse(self, file_path: str | Path, version_id: str) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ParserDependencyError("python-pptx is required for PPTX parsing") from exc

        presentation = Presentation(str(file_path))
        pages: list[PageIR] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            elements: list[Element] = []
            unsupported_count = 0
            for shape_index, shape in enumerate(slide.shapes, start=1):
                bbox = [
                    float(shape.left),
                    float(shape.top),
                    float(shape.left + shape.width),
                    float(shape.top + shape.height),
                ]
                if getattr(shape, "has_text_frame", False):
                    text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
                    if text:
                        elements.append(
                            Element(
                                id=f"{version_id}_slide_{slide_index:03d}_el_{shape_index:03d}",
                                type=ElementType.TEXT,
                                bbox=bbox,
                                order=shape_index,
                                source="native",
                                confidence=1.0,
                                text=text,
                            )
                        )
                elif getattr(shape, "has_table", False):
                    table = shape.table
                    rows = [[cell.text for cell in row.cells] for row in table.rows]
                    headers = rows[0] if rows else []
                    elements.append(
                        Element(
                            id=f"{version_id}_slide_{slide_index:03d}_table_{shape_index:03d}",
                            type=ElementType.TABLE,
                            bbox=bbox,
                            order=shape_index,
                            source="native",
                            confidence=1.0,
                            headers=headers,
                            rows=rows[1:] if len(rows) > 1 else [],
                            quality_flags=["TABLE_PRESENT"],
                        )
                    )
                elif shape.shape_type == 13:
                    elements.append(
                        Element(
                            id=f"{version_id}_slide_{slide_index:03d}_image_{shape_index:03d}",
                            type=ElementType.IMAGE,
                            bbox=bbox,
                            order=shape_index,
                            source="native",
                            confidence=1.0,
                            quality_flags=["VISUAL_CONTENT"],
                        )
                    )
                else:
                    unsupported_count += 1

            raw_text = "\n".join(element.text or "" for element in elements if element.text)
            flags = detect_page_quality(
                raw_text=raw_text,
                element_count=len(elements),
                image_count=sum(1 for element in elements if element.type == ElementType.IMAGE),
                unsupported_count=unsupported_count,
            )
            pages.append(
                PageIR(
                    id=f"{version_id}_slide_{slide_index:03d}",
                    version_id=version_id,
                    locator=Locator(type="slide", number=slide_index),
                    width=float(presentation.slide_width),
                    height=float(presentation.slide_height),
                    raw_text=raw_text,
                    elements=elements,
                    extraction_method="pptx_native",
                    quality_flags=flags,
                )
            )
        return ParsedDocument(version_id=version_id, pages=pages, native_result=pages, selected_result=pages)
